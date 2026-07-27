"""
替换 & 直接上传路由：支持手动修改 Word 后替换、跳过处理直接入库。
"""

from pathlib import Path

from fastapi import APIRouter, UploadFile, File
from fastapi.responses import JSONResponse

from web_app import state_manager
from web_app.config import (
    UPLOAD_DIR,
    OUTPUT_DIR,
    ALLOWED_EXTENSIONS,
    MAX_FILE_SIZE_MB,
)

router = APIRouter(prefix="/api", tags=["replace"])


def _count_chunks(docx_path: Path) -> int:
    """统计 DOCX 中 *** 标记的数量，估算 chunk 数。"""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(docx_path))
        count = 0
        for para in doc.paragraphs:
            if "***" in para.text:
                count += 1
        return count
    except Exception:
        return 0


def _extract_preview_text(docx_path: Path, max_lines: int = 200) -> str:
    """从 DOCX 提取纯文本用于预览。"""
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(docx_path))
        lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(text)
            if len(lines) >= max_lines:
                lines.append("... (预览截断，完整内容请下载查看)")
                break
        return "\n".join(lines)
    except Exception as e:
        return f"[预览提取失败: {e}]"


@router.post("/jobs/{job_id}/replace")
async def replace_output(job_id: str, file: UploadFile = File(...)):
    """上传手动编辑后的 DOCX 替换指定任务的输出文件。

    允许在 completed / ingested 状态下替换。
    替换后 ingested → completed，可重新入库。
    """
    # 1. 校验 job 存在且状态合法
    job = state_manager.get_job(job_id)
    if not job:
        return JSONResponse({"error": "任务不存在"}, status_code=404)
    if job["status"] not in ("completed", "ingested"):
        return JSONResponse(
            {"error": f"任务状态为 {job['status']}，仅已完成/已入库的任务可替换"},
            status_code=400,
        )

    # 2. 校验上传文件
    if not file.filename:
        return JSONResponse({"error": "未选择文件"}, status_code=400)

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        return JSONResponse({"error": f"不支持的文件格式: {ext}，仅支持 .docx/.doc"}, status_code=400)

    # 3. 读取文件内容，校验大小
    content = await file.read()
    file_size = len(content)
    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if file_size > max_bytes:
        return JSONResponse(
            {"error": f"文件过大: {file_size / 1024 / 1024:.1f}MB，最大 {MAX_FILE_SIZE_MB}MB"},
            status_code=400,
        )

    # 4. 保存修改版文件
    job_output_dir = OUTPUT_DIR / job_id
    job_output_dir.mkdir(parents=True, exist_ok=True)

    modify_ver = job.get("modify_count", 0) + 1
    new_filename = f"modified_v{modify_ver}_{file.filename}"
    new_path = job_output_dir / new_filename
    new_path.write_bytes(content)

    # 5. 提取预览和分块统计
    try:
        preview_text = _extract_preview_text(new_path)
        chunk_count = _count_chunks(new_path)
    except Exception as e:
        # 文件损坏或无法读取，删除已保存的文件
        if new_path.exists():
            new_path.unlink()
        return JSONResponse({"error": f"无法读取 DOCX 文件: {e}"}, status_code=500)

    # 6. 如果原任务已入库，先清理向量库中的旧切片
    was_ingested = job["status"] == "ingested"
    vector_cleaned = False
    if was_ingested:
        try:
            from sqlalchemy import create_engine, text
            import os as _os
            db_conn = _os.getenv(
                "PGVECTOR_CONNECTION",
                "postgresql+psycopg2://postgres:123456@10.32.10.161:5432/text2sql_vector",
            )
            engine = create_engine(db_conn)
            with engine.connect() as conn:
                result = conn.execute(
                    text("DELETE FROM parent_child_db_1024 WHERE c_metadata->>'source' = :src"),
                    {"src": job_id},
                )
                conn.commit()
                deleted = result.rowcount
            engine.dispose()
            vector_cleaned = True
            print(f"[replace] 已清理向量库中 {deleted} 条旧切片 (source={job_id})")
        except Exception as e:
            print(f"[replace] 清理向量库失败（将跳过）: {e}")

    # 7. 更新任务状态
    updated_job = state_manager.replace_job_output(
        job_id=job_id,
        output_path=str(new_path),
        preview_text=preview_text,
        chunk_count=chunk_count,
    )

    if updated_job is None:
        return JSONResponse({"error": "更新任务状态失败"}, status_code=500)

    msg = "文件已替换，可重新入库"
    if vector_cleaned:
        msg += "（已清理旧向量）"

    return JSONResponse({
        "status": "ok",
        "message": msg,
        "job_id": job_id,
        "chunk_count": chunk_count,
        "job": updated_job,
    })


@router.post("/direct-upload")
async def direct_upload(file: UploadFile = File(...)):
    """直接上传手工编辑好的 DOCX，跳过 MinerU 处理，创建 completed 状态任务。

    适用于：用户已有手工编辑好的带分块标记的 DOCX，只需入库的场景。
    """
    # 1. 校验文件
    if not file.filename:
        return JSONResponse({"error": "未选择文件"}, status_code=400)

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        return JSONResponse({"error": f"不支持的文件格式: {ext}，仅支持 .docx/.doc"}, status_code=400)

    content = await file.read()
    file_size = len(content)
    max_bytes = MAX_FILE_SIZE_MB * 1024 * 1024
    if file_size > max_bytes:
        return JSONResponse(
            {"error": f"文件过大: {file_size / 1024 / 1024:.1f}MB，最大 {MAX_FILE_SIZE_MB}MB"},
            status_code=400,
        )

    # 2. 创建任务
    job_id = state_manager.create_job(
        filename=file.filename,
        file_size=file_size,
        input_path="",
    )

    # 3. 保存文件到 output 目录
    job_output_dir = OUTPUT_DIR / job_id
    job_output_dir.mkdir(parents=True, exist_ok=True)
    output_path = job_output_dir / file.filename
    output_path.write_bytes(content)

    # 4. 提取预览和分块统计
    try:
        preview_text = _extract_preview_text(output_path)
        chunk_count = _count_chunks(output_path)
    except Exception as e:
        state_manager.delete_job(job_id)
        return JSONResponse({"error": f"无法读取 DOCX 文件: {e}"}, status_code=500)

    # 5. 更新为 completed 状态（跳过 MinerU 处理）
    state_manager.update_job(
        job_id,
        status="completed",
        output_path=str(output_path),
        progress="已就绪（跳过处理，可直接入库）",
        chunk_count=chunk_count,
        preview_text=preview_text,
    )

    return JSONResponse({
        "job_id": job_id,
        "status": "completed",
        "filename": file.filename,
        "chunk_count": chunk_count,
        "message": "文档已就绪，可直接入库",
    })
