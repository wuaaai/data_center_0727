"""
RAG 入库脚本：将 DOCX 切片并写入 pgvector 向量库。
从 Langchain_160 迁移至 mineru 项目，路径全部改为基于 mineru 项目根目录。
"""
import os
import json
import hashlib
import requests
import sys
import time
from pathlib import Path
from typing import List

from docx import Document as DocxDocument
from langchain_postgres import PGVector
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings

# ---------- 项目根目录 ----------
_PROJECT_ROOT = Path(__file__).resolve().parent.parent  # data_processing_center/
_WEB_APP_DIR = Path(__file__).resolve().parent           # mineru_app/

# ---------- 配置 ----------
COLLECTION_NAME = os.getenv("PGVECTOR_COLLECTION_NAME", "parent_child_db_1024")
DB_CONNECTION = os.getenv(
    "PGVECTOR_CONNECTION",
    "postgresql+psycopg2://postgres:123456@10.32.10.161:5432/text2sql_vector",
)
LOCAL_API_URL = os.getenv("EMBEDDING_API_URL", "http://120.211.116.133:65325/embed")
EMBEDDING_TOKEN = "HB|Ch0.1415926"
IMAGE_SAVE_DIR = _PROJECT_ROOT / "output" / "images"
IMAGE_BASE_URL = "http://localhost:8001/static/images/"

# 区划映射表
MAPPING_PATH = _WEB_APP_DIR / "region_mapping.json"
region_mapping = {}
if MAPPING_PATH.exists():
    with open(MAPPING_PATH, "r", encoding="utf-8") as f:
        region_mapping = json.load(f)
    print(f"[region] 已加载区划映射: {len(region_mapping)} 个文件")
else:
    print("[WARNING] 未找到 region_mapping.json，入库时 region_code 将为默认值")


# ---------- Embedding 模型 ----------
class CustomLocalEmbeddings(Embeddings):
    def __init__(self, api_url: str):
        self.api_url = api_url

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        try:
            response = requests.post(
                self.api_url,
                params={"token": EMBEDDING_TOKEN},
                json=texts,
                headers={"Content-Type": "application/json", "accept": "application/json"},
                timeout=60,
            )
            response.raise_for_status()
            return response.json()["embeddings"]
        except Exception as e:
            print(f"[ERROR] 请求 Embedding 接口失败: {e}")
            return []

    def embed_query(self, text: str) -> list[float]:
        embeddings = self.embed_documents([text])
        if embeddings and len(embeddings) > 0:
            return embeddings[0]
        return []


print(f"[LOADING] 加载自定义接口模型: {LOCAL_API_URL} ...")
EMBEDDING_MODEL = CustomLocalEmbeddings(api_url=LOCAL_API_URL)

IMAGE_SAVE_DIR.mkdir(parents=True, exist_ok=True)


# ---------- 图片提取 ----------
def extract_images_from_paragraph(paragraph):
    md_links = []
    for run in paragraph.runs:
        if 'drawing' in run._element.xml:
            blips = run._element.xpath('.//a:blip')
            for blip in blips:
                embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if embed_id:
                    image_part = paragraph.part.related_parts[embed_id]
                    image_bytes = image_part.blob
                    img_hash = hashlib.md5(image_bytes).hexdigest()
                    content_type = image_part.content_type
                    ext = content_type.split('/')[-1] if '/' in content_type else 'png'
                    if ext == 'jpeg':
                        ext = 'jpg'
                    filename = f"{img_hash}.{ext}"
                    save_path = IMAGE_SAVE_DIR / filename
                    if not save_path.exists():
                        save_path.write_bytes(image_bytes)
                    img_url = f"{IMAGE_BASE_URL}{filename}"
                    md_links.append(f"\n![image]({img_url})\n")
    return "".join(md_links)


# ---------- 文档处理 ----------
def read_and_process_word(file_path, source_id=None, original_filename=None):
    try:
        doc = DocxDocument(file_path)
        file_name = os.path.basename(file_path)
        if source_id is None:
            source_id = file_name
        full_text_list = []
        for para in doc.paragraphs:
            text = para.text.strip()
            img_md = extract_images_from_paragraph(para)
            if text:
                full_text_list.append(text)
            if img_md:
                full_text_list.append(img_md)

        full_text_with_images = "\n".join(full_text_list)
        docs_to_save = []
        parent_blocks = full_text_with_images.split("***")

        for p_idx, parent_block in enumerate(parent_blocks):
            if not parent_block.strip():
                continue
            parent_title = ""
            lines = parent_block.strip().split('\n')
            for line in lines:
                if "文档名：" in line:
                    parent_title = line.strip()
                    break
            if not parent_title:
                parent_title = f"文档名：{file_name}"

            full_parent_context = parent_block.strip()
            child_blocks = parent_block.split("<-split->")

            for c_idx, child_content in enumerate(child_blocks):
                child_content = child_content.strip()
                if not child_content:
                    continue
                content_to_vectorize = f"{parent_title}\n{child_content}"
                metadata = {
                    "source": source_id,
                    "filename": original_filename,
                    "type": "child_chunk",
                    "recall_context": full_parent_context,
                    "chunk_id": f"{file_name}_p{p_idx}_c{c_idx}",
                    "region_code": region_mapping.get(file_name, "000000"),
                    "created_at": time.time(),
                }
                new_doc = Document(page_content=content_to_vectorize, metadata=metadata)
                docs_to_save.append(new_doc)

        return docs_to_save

    except Exception as e:
        print(f" 解析文件 {file_path} 失败: {e}")
        import traceback
        traceback.print_exc()
        return []


# ---------- 入库 ----------
def ingest(file_path, source_id=None):
    if not os.path.exists(file_path):
        print(f"文件不存在: {file_path}")
        return

    file_name = os.path.basename(file_path)
    if source_id is None:
        source_id = file_name

    # 从 job 记录中获取原始中文文档名
    original_filename = file_name
    try:
        import json as _json
        jobs_file = _PROJECT_ROOT / "output" / "_jobs.json"
        if jobs_file.exists():
            with open(jobs_file, "r", encoding="utf-8") as f:
                jobs = _json.load(f)
            job = jobs.get(source_id)
            if job and job.get("filename"):
                original_filename = job["filename"]
    except Exception:
        pass
    print(f"  原始文档名: {original_filename}")

    print(f"连接 PostgreSQL 数据库: {COLLECTION_NAME} (维度: 1024)")
    # 解析 DB_CONNECTION 为 psycopg2 格式
    import psycopg2
    db_url = DB_CONNECTION
    # postgresql+psycopg2://postgres:123456@10.32.10.161:5432/text2sql_vector
    # -> host=10.32.10.161 port=5432 user=postgres password=123456 dbname=text2sql_vector
    from urllib.parse import urlparse
    parsed = urlparse(db_url.replace("postgresql+psycopg2://", "postgresql://"))
    conn = psycopg2.connect(
        host=parsed.hostname, port=parsed.port or 5432,
        user=parsed.username, password=parsed.password,
        dbname=parsed.path.lstrip("/")
    )
    cur = conn.cursor()

    # 1) 按 source_id 删除旧切片
    try:
        cur.execute("DELETE FROM parent_child_db_1024 WHERE c_metadata->>'source' = %s", (source_id,))
        deleted = cur.rowcount
        conn.commit()
        print(f"已清理旧切片 (source={source_id}, 删除{deleted}条)")
    except Exception as e:
        conn.rollback()
        print(f"清理旧切片跳过: {e}")

    docs = read_and_process_word(file_path, source_id=source_id, original_filename=original_filename)
    print(f"  - {file_name}: 生成 {len(docs)} 个切片")

    if docs:
        print(f"正在写入 {len(docs)} 个切片到PostgreSQL数据库...")
        texts = [d.page_content for d in docs]
        metadatas = [d.metadata for d in docs]

        # 调用 embedding API（逐个发送，避免批量超长导致 500）
        print(f"请求 Embedding API: {LOCAL_API_URL}")
        embeddings = []
        for i, text in enumerate(texts):
            if i % 3 == 0:
                print(f"  embedding 进度: {i+1}/{len(texts)}")
            resp = requests.post(LOCAL_API_URL, params={"token": EMBEDDING_TOKEN}, json=[text],
                                 headers={"Content-Type": "application/json", "accept": "application/json"}, timeout=60)
            resp.raise_for_status()
            emb_list = resp.json()["embeddings"]
            if emb_list:
                embeddings.append(emb_list[0])
        print(f"获取到 {len(embeddings)} 个向量 (维度: {len(embeddings[0]) if embeddings else 'N/A'})")

        for text, meta, emb in zip(texts, metadatas, embeddings):
            cur.execute(
                "INSERT INTO parent_child_db_1024 (c_document, c_embedding, c_metadata) VALUES (%s, %s::vector, %s)",
                (text, str(emb), json.dumps(meta, ensure_ascii=False))
            )
        conn.commit()
        print("[OK] 入库完成！")
    else:
        print("无数据入库")
    cur.close()
    conn.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python rag_ingest.py <file_path> [source_id]")
        print("示例: python rag_ingest.py D:/path/to/file.docx abc123")
        sys.exit(1)
    src_id = sys.argv[2] if len(sys.argv) > 2 else None
    ingest(sys.argv[1], source_id=src_id)
